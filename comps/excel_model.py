"""Modelo de valuation em Excel, no layout do template do Matheus (modelo da Alphabet).

O que vem do template: fonte Avenir Next LT Pro, paleta vermelha (faixas B1141A / EC1B23 /
CB1720, cabeçalhos de tabela FF7C80), números em cinza F2F2F2, legenda Hardcode / Fórmula /
Premissa no topo de cada aba, nome da empresa e ticker no canto superior esquerdo, a aba
"Cenários" com seletor 1/2/3 e linha "Selecionado" via CHOOSE, e a aba DCF com o resumo à
esquerda (preço-alvo, upside) e o fluxo por ano à direita.

Convenções de cor nas células: azul = hardcode (dado da CVM/mercado), azul em fundo amarelo
= premissa editável, preto = fórmula, verde = link de outra aba.
Tudo que é cálculo é fórmula: mudar o seletor de cenário, uma premissa ou o beta recalcula
DCF, recomendação, sensibilidade e football field.
"""
from __future__ import annotations

import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .dcf import SCENARIOS, YEARS
from .valuation import BUY_UPSIDE, SELL_UPSIDE, ValuationResult

FONT = "Avenir Next LT Pro"
DARK, RED, RED2, PINK = "B1141A", "EC1B23", "CB1720", "FF7C80"
GRAY, YELLOW, TARGET = "F2F2F2", "FFEB9C", "FFFF00"
HARD, LINK, BLACK, WHITE, MUTED = "0000FF", "008000", "000000", "FFFFFF", "7F7F7F"

F_MM = '#,##0;(#,##0);"-"'
F_PCT = '0.0%;(0.0%);"-"'
F_PX = '"R$ "#,##0.00'
F_X = '0.0"x"'
F_SH = '#,##0.0'

THIN = Side(style="thin", color="BFBFBF")
MEDIUM = Side(style="medium", color=BLACK)
DOUBLE = Side(style="double", color=BLACK)
DASHED = Side(style="dashed", color=BLACK)

SC_PT = {"pessimista": "Pessimista", "base": "Base", "otimista": "Otimista"}
YRS = ["D", "E", "F", "G", "H"]          # Ano 1..5 nas abas Cenários e DCF Cenários
DCF_YRS = ["K", "L", "M", "N", "O"]      # Ano 1..5 na aba DCF


# ---------------------------------------------------------------------- estilo
def _font(color=BLACK, bold=False, size=11, italic=False):
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic)


def _put(ws, ref, value, fmt=None, kind="formula", bold=False, gray=False, size=11, align=None):
    """kind: formula (preto) | hard (azul) | assump (azul em amarelo) | link (verde) | text."""
    c = ws[ref]
    c.value = value
    color = {"formula": BLACK, "hard": HARD, "assump": HARD, "link": LINK, "text": BLACK}[kind]
    c.font = _font(color, bold=bold, size=size)
    if kind == "assump":
        c.fill = PatternFill("solid", fgColor=YELLOW)
    elif gray:
        c.fill = PatternFill("solid", fgColor=GRAY)
    if fmt:
        c.number_format = fmt
    if align or fmt:
        c.alignment = Alignment(horizontal=align or "right")
    return c


def _band(ws, first: str, last_col: str, text: str, fill=DARK, size=12, bold=True):
    """Faixa colorida estilo template: preenche da coluna de `first` até `last_col`, texto na primeira."""
    col0 = ws[first].column
    row = ws[first].row
    col1 = ws[f"{last_col}{row}"].column
    for c in range(col0, col1 + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = _font(WHITE, bold=bold, size=size)
    ws[first].value = text


def _label(ws, ref, text, bold=False, size=11, color=BLACK, italic=False):
    c = ws[ref]
    c.value = text
    c.font = _font(color, bold=bold, size=size, italic=italic)
    return c


def _frame(ws, title: tuple[str, str], legend_col: str, widths: dict[str, float]):
    """Cabeçalho do template, sem imagens: nome da empresa e ticker (linhas 1-3) e legenda de cores."""
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 18
    _label(ws, "B1", title[0], bold=True, size=16, color=DARK)
    _label(ws, "B2", title[1], size=10, color=MUTED)
    _label(ws, f"{legend_col}1", "Hardcode", color="0070C0", size=10)
    _label(ws, f"{legend_col}2", "Fórmula", size=10)
    c = _label(ws, f"{legend_col}3", "Premissa", color="FF3F3F", size=10)
    c.fill = PatternFill("solid", fgColor="FFFA93")


def _rec_format(ws, ref: str):
    green = PatternFill("solid", bgColor="C6EFCE")
    red = PatternFill("solid", bgColor="FFC7CE")
    grey = PatternFill("solid", bgColor="E7E6E6")
    ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"COMPRA"'], fill=green,
                                                  font=Font(color="006100", bold=True)))
    ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"VENDA"'], fill=red,
                                                  font=Font(color="9C0006", bold=True)))
    ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"NEUTRO"'], fill=grey,
                                                  font=Font(color="404040", bold=True)))


def _num(v):
    return None if v is None or (isinstance(v, float) and math.isnan(v)) else float(v)


# ---------------------------------------------------------------------- export
def export_model(v: ValuationResult, path: Path | str, comps=None) -> Path:
    """`comps` = CompsResult (para a aba Comps). Sem ele, a aba usa só as faixas já calculadas."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    short = v.company.title().replace(" S.A.", "").replace(" S.a.", "")
    header = (short, f"{v.ticker} | Modelo de valuation")

    wb = Workbook()
    ws_r = wb.active
    ws_r.title = "Resumo"
    ws_c = wb.create_sheet("Cenários")
    ws_d = wb.create_sheet("DCF")
    ws_dc = wb.create_sheet("DCF Cenários")
    ws_cp = wb.create_sheet("Comps")
    ws_s = wb.create_sheet("Sensibilidade")
    ws_t = wb.create_sheet("Teses")
    ws_h = wb.create_sheet("Histórico")

    # ================================================================== Cenários
    ws = ws_c
    _frame(ws, header, "L", {"B": 16, "C": 8, **{c: 10 for c in YRS}, "I": 2, "J": 70})
    _label(ws, "B5", "Cenário", bold=True)
    _put(ws, "C5", 2, "0", "assump", align="center")
    _label(ws, "D5", "1 = Pessimista · 2 = Base · 3 = Otimista", color=MUTED, size=9, italic=True)
    _put(ws, "B6", '=IF($C$5=2,"Cenário base",IF($C$5=1,"Cenário pessimista","Cenário otimista"))', kind="formula",
         bold=True, align="left")
    ws["B6"].border = Border(bottom=THIN)
    drivers = [("Crescimento da receita (%)", "growth", True, "crescimento"),
               ("Margem EBITDA ex-IFRS16 (%)", "margin", True, "margem"),
               ("Capex / receita (%)", "capex_pct", False, "capex"),
               ("D&A sem direito de uso / receita (%)", "da_pct", False, None),
               ("Capital de giro / receita (%)", "nwc_pct", False, "capital de giro")]
    SEL, SROW = {}, {}
    r = 8
    for title, key, multi, rkey in drivers:
        _band(ws, f"D{r}", "H", title, fill=RED2, size=11)
        if multi:
            for i, c in enumerate(YRS):
                _label(ws, f"{c}{r + 1}", f"Ano {i + 1}")
                ws[f"{c}{r + 1}"].alignment = Alignment(horizontal="right")
        else:
            _label(ws, f"D{r + 1}", "Todos os anos")
        _label(ws, f"J{r}", "Justificativa (ligada às teses)", bold=True, size=10)
        _label(ws, f"B{r + 2}", "Cenário")
        _put(ws, f"C{r + 2}", "=$C$5", "0", "formula", align="center")
        SROW[key] = {}
        for j, k in enumerate(SCENARIOS):
            rr = r + 3 + j
            SROW[key][k] = rr
            cell = _label(ws, f"B{rr}", SC_PT[k], bold=True, color=WHITE)
            cell.fill = PatternFill("solid", fgColor=RED2)
            sc = v.scenarios.get(k)
            if sc is None:
                continue
            vals = getattr(sc, key)
            if multi:
                for c, x in zip(YRS, vals):
                    _put(ws, f"{c}{rr}", x, F_PCT, "assump")
            else:
                _put(ws, f"D{rr}", vals, F_PCT, "assump")
            if rkey and sc.rationale.get(rkey):
                t = ws[f"J{rr}"]
                t.value = sc.rationale[rkey]
                t.font = _font(MUTED, size=9)
                t.alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[rr].height = 30 if len(sc.rationale[rkey]) > 90 else 15
        rr = r + 6
        SEL[key] = rr
        cell = _label(ws, f"B{rr}", "Selecionado", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=RED2)
        for c in (YRS if multi else ["D"]):
            _put(ws, f"{c}{rr}", f"=CHOOSE($C${r + 2},{c}{r + 3},{c}{r + 4},{c}{r + 5})", F_PCT, "formula")
        for c in ["B", "C"] + (YRS if multi else ["D"]):
            ws[f"{c}{rr}"].border = Border(bottom=MEDIUM)
        r += 9
    used = [f"{SC_PT[k]}: " + ", ".join(f"#{i}" for i in sc.theses_used) for k, sc in v.scenarios.items() if sc.theses_used]
    if used:
        _label(ws, f"B{r}", "Teses usadas por cenário — " + " · ".join(used), color=MUTED, size=9, italic=True)
    _label(ws, f"B{r + 1}", f"Fonte das premissas: {v.assumptions_source}", color=MUTED, size=9, italic=True)
    ws.freeze_panes = "A7"

    # ================================================================== DCF (cenário selecionado)
    ws = ws_d
    _frame(ws, header, "R", {"B": 2, "C": 2, "D": 40, "E": 2, "F": 2, "G": 14, "H": 3, "I": 40, "J": 12,
                          **{c: 11 for c in DCF_YRS}, "P": 15, "Q": 2, "R": 12})
    b = v.base
    _label(ws, "D5", "Data de fechamento")
    c = _put(ws, "G5", "=TODAY()", "dd/mm/yyyy", "formula")
    c.font = _font("FF0000")
    for col in ("D", "G"):
        ws[f"{col}5"].border = Border(bottom=THIN)
    _band(ws, "I6", "I", "DCF (FCFF) — R$ mi")
    _band(ws, "J6", "J", "LTM", fill=RED)
    for i, col in enumerate(DCF_YRS):
        _band(ws, f"{col}6", col, f"Ano {i + 1}", fill=RED)
    _band(ws, "P6", "P", "Perpetuidade", fill=RED)
    for col in ["J"] + DCF_YRS:
        ws[f"{col}6"].alignment = Alignment(horizontal="right")
    ws["P6"].alignment = Alignment(horizontal="center")

    left = [
        (7, "Preço atual (R$)", v.price, F_PX, "hard"),
        (8, "Ações em circulação (mi)", v.shares / 1e6, F_SH, "hard"),
        (9, "(−) Dívida líquida ex-arrendamento (R$ mi)", v.wacc.net_debt, F_MM, "hard"),
        (10, "(−) Minoritários (R$ mi)", v.minorities, F_MM, "hard"),
    ]
    for row, label, val, fmt, kind in left:
        _label(ws, f"D{row}", label)
        _put(ws, f"G{row}", val, fmt, kind)
    _band(ws, "D12", "G", "PREÇO-ALVO (FCFF)", fill=RED)
    _label(ws, "D13", "Preço-alvo (R$)", bold=True)
    c = _put(ws, "G13", "=G18/G8", F_PX, "formula", bold=True)
    c.fill = PatternFill("solid", fgColor=TARGET)
    c.border = Border(bottom=THIN)
    _label(ws, "D14", "Upside (%)")
    _put(ws, "G14", "=G13/G7-1", F_PCT, "formula", bold=True)
    _label(ws, "D15", "Recomendação", bold=True)
    _put(ws, "G15", '=IF(G14>=G22,"COMPRA",IF(G14<=G23,"VENDA","NEUTRO"))', kind="formula", bold=True, align="center")
    _rec_format(ws, "G15")
    _label(ws, "D17", "Enterprise value (R$ mi)", bold=True)
    _put(ws, "G17", "=SUM(K22:O22)+P22", F_MM, "formula", bold=True)
    for col in ("D", "G"):
        ws[f"{col}17"].border = Border(bottom=DASHED)
    _label(ws, "D18", "Equity value (R$ mi)", bold=True)
    _put(ws, "G18", "=G17-G9-G10", F_MM, "formula", bold=True)
    _band(ws, "D21", "G", "CRITÉRIO DE RECOMENDAÇÃO", fill=RED)
    _label(ws, "D22", "COMPRA se upside ≥")
    _put(ws, "G22", BUY_UPSIDE, F_PCT, "assump")
    _label(ws, "D23", "VENDA se upside ≤")
    _put(ws, "G23", SELL_UPSIDE, F_PCT, "assump")
    _label(ws, "D24", "Entre os dois: NEUTRO", color=MUTED, size=9, italic=True)
    _label(ws, "D26", "% do EV vindo da perpetuidade")
    _put(ws, "G26", "=P22/G17", F_PCT, "formula")
    _label(ws, "D27", "Cenário em uso")
    _put(ws, "G27", "='Cenários'!B6", kind="link", align="right")

    S = lambda key, col: f"'Cenários'!{col}{SEL[key]}"                           # noqa: E731
    S1 = lambda key: f"'Cenários'!$D${SEL[key]}"                                  # noqa: E731
    rows = {"rec": 7, "cres": 8, "marg": 9, "ebitda": 10, "da": 11, "ebit": 12, "ir": 13, "nopat": 14, "da_add": 15,
            "capex": 16, "cg": 17, "dcg": 18, "fcff": 19, "per": 20, "fat": 21, "vp": 22}
    labels = {"rec": "Receita líquida", "cres": "% a/a", "marg": "Margem EBITDA ex-IFRS16", "ebitda": "EBITDA ex-IFRS16",
              "da": "(−) D&A sem direito de uso", "ebit": "EBIT", "ir": "(−) IR/CS sobre o EBIT", "nopat": "NOPAT",
              "da_add": "(+) D&A", "capex": "(−) Capex", "cg": "Capital de giro", "dcg": "(+/−) Δ capital de giro",
              "fcff": "Fluxo de caixa livre da firma (FCFF)", "per": "Período (anos)", "fat": "Fator de desconto",
              "vp": "DCF (FCFF)"}
    for key, row in rows.items():
        _label(ws, f"I{row}", labels[key], bold=key in ("rec", "ebitda", "ebit", "nopat", "fcff", "vp"),
               size=11 if key not in ("cres", "marg") else 10)
    R = rows
    _put(ws, f"J{R['rec']}", b.revenue, F_MM, "hard", gray=True)
    _put(ws, f"J{R['ebitda']}", b.ebitda_ex16, F_MM, "hard", gray=True)
    _put(ws, f"J{R['marg']}", f"=J{R['ebitda']}/J{R['rec']}", F_PCT, "formula", gray=True)
    _put(ws, f"J{R['cg']}", f"=J{R['rec']}*{S1('nwc_pct')}", F_MM, "link", gray=True)
    for i, c in enumerate(DCF_YRS):
        prev = "J" if i == 0 else DCF_YRS[i - 1]
        sc_col = YRS[i]
        _put(ws, f"{c}{R['cres']}", f"={S('growth', sc_col)}", F_PCT, "link", gray=True)
        _put(ws, f"{c}{R['rec']}", f"={prev}{R['rec']}*(1+{c}{R['cres']})", F_MM, gray=True, bold=True)
        _put(ws, f"{c}{R['marg']}", f"={S('margin', sc_col)}", F_PCT, "link", gray=True)
        _put(ws, f"{c}{R['ebitda']}", f"={c}{R['rec']}*{c}{R['marg']}", F_MM, gray=True)
        _put(ws, f"{c}{R['da']}", f"=-{c}{R['rec']}*{S1('da_pct')}", F_MM, "link", gray=True)
        _put(ws, f"{c}{R['ebit']}", f"={c}{R['ebitda']}+{c}{R['da']}", F_MM, gray=True)
        _put(ws, f"{c}{R['ir']}", f"=-MAX({c}{R['ebit']},0)*$J$33", F_MM, gray=True)
        _put(ws, f"{c}{R['nopat']}", f"={c}{R['ebit']}+{c}{R['ir']}", F_MM, gray=True)
        _put(ws, f"{c}{R['da_add']}", f"=-{c}{R['da']}", F_MM, gray=True)
        _put(ws, f"{c}{R['capex']}", f"=-{c}{R['rec']}*{S1('capex_pct')}", F_MM, "link", gray=True)
        _put(ws, f"{c}{R['cg']}", f"={c}{R['rec']}*{S1('nwc_pct')}", F_MM, "link", gray=True)
        _put(ws, f"{c}{R['dcg']}", f"=-({c}{R['cg']}-{prev}{R['cg']})", F_MM, gray=True)
        _put(ws, f"{c}{R['fcff']}", f"={c}{R['nopat']}+{c}{R['da_add']}+{c}{R['capex']}+{c}{R['dcg']}", F_MM,
             gray=True, bold=True)
        _put(ws, f"{c}{R['per']}", i + 1, "0", "hard", gray=True)
        _put(ws, f"{c}{R['fat']}", f"=1/(1+$J$38)^{c}{R['per']}", "0.000", gray=True)
        _put(ws, f"{c}{R['vp']}", f"={c}{R['fcff']}*{c}{R['fat']}", F_MM, gray=True, bold=True)
    _put(ws, f"P{R['fcff']}", f"=IF($J$38>$J$39,O{R['fcff']}*(1+$J$39)/($J$38-$J$39),NA())", F_MM, gray=True, bold=True)
    _put(ws, f"P{R['per']}", f"=O{R['per']}", "0", gray=True)
    _put(ws, f"P{R['fat']}", f"=O{R['fat']}", "0.000", gray=True)
    _put(ws, f"P{R['vp']}", f"=P{R['fcff']}*P{R['fat']}", F_MM, gray=True, bold=True)
    for c in ["I", "J"] + DCF_YRS + ["P"]:
        ws[f"{c}{R['fcff']}"].border = Border(top=THIN)
        ws[f"{c}{R['vp']}"].border = Border(bottom=DOUBLE)

    # WACC (mesmo bloco da aba DCF do template)
    _band(ws, "I27", "J", "CÁLCULO DO WACC")
    w = v.wacc
    wacc_rows = [
        (28, "Taxa livre de risco (R$, nominal)", w.rf, F_PCT, "assump", "Média da Selic esperada no Focus (23/10/2026) para o fim de 2026–2029"),
        (29, "Beta", w.beta, "0.00", "assump", w.beta_source),
        (30, "Prêmio de risco de mercado", w.erp, F_PCT, "assump", "Damodaran, prêmio implícito (jan/2026); risco-país já está na taxa em reais"),
        (31, "Custo do capital próprio (Ke)", "=J28+J29*J30", F_PCT, "formula", "CAPM"),
        (32, "Custo da dívida antes do IR", w.kd_pre, F_PCT, "assump", "Taxa livre de risco + spread de crédito"),
        (33, "IR/CS", w.tax, F_PCT, "assump", "Alíquota nominal (IR 25% + CSLL 9%)"),
        (34, "Custo da dívida depois do IR (Kd)", "=J32*(1-J33)", F_PCT, "formula", ""),
        (35, "Valor de mercado (R$ mi)", w.equity_value, F_MM, "hard", "Ações × preço por classe"),
        (36, "Dívida líquida ex-arrendamento (R$ mi)", "=G9", F_MM, "formula", ""),
        (37, "Peso da dívida D/(D+E)", "=MAX(J36,0)/(MAX(J36,0)+J35)", F_PCT, "formula", "Caixa líquido não vira peso negativo"),
        (38, "WACC", "=(1-J37)*J31+J37*J34", F_PCT, "formula", ""),
        (39, "Crescimento na perpetuidade (g)", v.dcf["base"].scenario.terminal_growth if "base" in v.dcf else 0.045,
         F_PCT, "assump", "Inflação de longo prazo (~3,7%) + ~1% real"),
    ]
    for row, label, val, fmt, kind, src in wacc_rows:
        _label(ws, f"I{row}", label, bold=label == "WACC")
        _put(ws, f"J{row}", val, fmt, kind, bold=label == "WACC")
        if src:
            _label(ws, f"K{row}", src, color=MUTED, size=9, italic=True)
    ws["J38"].border = Border(top=THIN, bottom=DOUBLE)

    # ================================================================== DCF Cenários (os três, para o football field)
    ws = ws_dc
    _frame(ws, header, "J", {"B": 40, "C": 12, **{c: 11 for c in YRS}})
    _label(ws, "B5", "DCF dos três cenários (mesma conta da aba DCF, sem depender do seletor)", bold=True)
    PS = {}
    r = 7
    for k in SCENARIOS:
        if k not in v.scenarios:
            continue
        _band(ws, f"B{r}", "B", f"Cenário {SC_PT[k]} (R$ mi)")
        _band(ws, f"C{r}", "C", "LTM", fill=RED)
        for i, c in enumerate(YRS):
            _band(ws, f"{c}{r}", c, f"Ano {i + 1}", fill=RED)
        names = ["per", "rec", "cres", "marg", "ebitda", "da", "ebit", "ir", "capex", "cg", "dcg", "fcff", "fat", "vp"]
        Q = {n: r + 1 + i for i, n in enumerate(names)}
        lab = {"per": "Período (anos)", "rec": "Receita", "cres": "Crescimento", "marg": "Margem EBITDA ex-IFRS16",
               "ebitda": "EBITDA ex-IFRS16", "da": "(−) D&A sem direito de uso", "ebit": "EBIT", "ir": "(−) IR/CS",
               "capex": "(−) Capex", "cg": "Capital de giro", "dcg": "(+/−) Δ capital de giro", "fcff": "FCFF",
               "fat": "Fator de desconto", "vp": "Valor presente do FCFF"}
        for n, rr in Q.items():
            _label(ws, f"B{rr}", lab[n], bold=n in ("rec", "fcff"))
        row_of = lambda key: SROW[key][k]                                         # noqa: E731
        _put(ws, f"C{Q['rec']}", "=DCF!$J$7", F_MM, "link", gray=True)
        _put(ws, f"C{Q['marg']}", "=DCF!$J$9", F_PCT, "link", gray=True)
        _put(ws, f"C{Q['cg']}", f"=C{Q['rec']}*'Cenários'!$D${row_of('nwc_pct')}", F_MM, "link", gray=True)
        for i, c in enumerate(YRS):
            prev = "C" if i == 0 else YRS[i - 1]
            _put(ws, f"{c}{Q['per']}", i + 1, "0", "hard", gray=True)
            _put(ws, f"{c}{Q['cres']}", f"='Cenários'!{c}{row_of('growth')}", F_PCT, "link", gray=True)
            _put(ws, f"{c}{Q['rec']}", f"={prev}{Q['rec']}*(1+{c}{Q['cres']})", F_MM, gray=True)
            _put(ws, f"{c}{Q['marg']}", f"='Cenários'!{c}{row_of('margin')}", F_PCT, "link", gray=True)
            _put(ws, f"{c}{Q['ebitda']}", f"={c}{Q['rec']}*{c}{Q['marg']}", F_MM, gray=True)
            _put(ws, f"{c}{Q['da']}", f"=-{c}{Q['rec']}*'Cenários'!$D${row_of('da_pct')}", F_MM, "link", gray=True)
            _put(ws, f"{c}{Q['ebit']}", f"={c}{Q['ebitda']}+{c}{Q['da']}", F_MM, gray=True)
            _put(ws, f"{c}{Q['ir']}", f"=-MAX({c}{Q['ebit']},0)*DCF!$J$33", F_MM, "link", gray=True)
            _put(ws, f"{c}{Q['capex']}", f"=-{c}{Q['rec']}*'Cenários'!$D${row_of('capex_pct')}", F_MM, "link", gray=True)
            _put(ws, f"{c}{Q['cg']}", f"={c}{Q['rec']}*'Cenários'!$D${row_of('nwc_pct')}", F_MM, "link", gray=True)
            _put(ws, f"{c}{Q['dcg']}", f"=-({c}{Q['cg']}-{prev}{Q['cg']})", F_MM, gray=True)
            _put(ws, f"{c}{Q['fcff']}", f"={c}{Q['ebitda']}+{c}{Q['ir']}+{c}{Q['capex']}+{c}{Q['dcg']}", F_MM,
                 gray=True, bold=True)
            _put(ws, f"{c}{Q['fat']}", f"=1/(1+DCF!$J$38)^{c}{Q['per']}", "0.000", "link", gray=True)
            _put(ws, f"{c}{Q['vp']}", f"={c}{Q['fcff']}*{c}{Q['fat']}", F_MM, gray=True)
        v0 = Q["vp"] + 2
        vals = [
            ("Soma do valor presente dos FCFF", f"=SUM(D{Q['vp']}:H{Q['vp']})", F_MM, "formula"),
            ("Valor terminal (Gordon)", f"=IF(DCF!$J$38>DCF!$J$39,H{Q['fcff']}*(1+DCF!$J$39)/(DCF!$J$38-DCF!$J$39),NA())", F_MM, "link"),
            ("Valor presente do valor terminal", f"=D{v0 + 1}*H{Q['fat']}", F_MM, "formula"),
            ("Enterprise value", f"=D{v0}+D{v0 + 2}", F_MM, "formula"),
            ("(−) Dívida líquida ex-arrendamento", "=-DCF!$G$9", F_MM, "link"),
            ("(−) Minoritários", "=-DCF!$G$10", F_MM, "link"),
            ("Equity value", f"=D{v0 + 3}+D{v0 + 4}+D{v0 + 5}", F_MM, "formula"),
            ("Valor por ação (R$)", f"=D{v0 + 6}/DCF!$G$8", F_PX, "link"),
            ("Upside", f"=D{v0 + 7}/DCF!$G$7-1", F_PCT, "link"),
            ("Recomendação", f'=IF(D{v0 + 8}>=DCF!$G$22,"COMPRA",IF(D{v0 + 8}<=DCF!$G$23,"VENDA","NEUTRO"))', None, "link"),
        ]
        for i, (lab_, fx, fmt, kind) in enumerate(vals):
            rr = v0 + i
            bold = lab_.startswith(("Enterprise", "Equity", "Valor por", "Recomendação"))
            _label(ws, f"B{rr}", lab_, bold=bold)
            _put(ws, f"D{rr}", fx, fmt, kind, bold=bold, align="right")
        _rec_format(ws, f"D{v0 + 9}")
        PS[k] = v0 + 7
        r = v0 + len(vals) + 2

    # ================================================================== Comps
    ws = ws_cp
    _frame(ws, header, "K", {"B": 30, "C": 9, "D": 11, "E": 11, "F": 11, "G": 13, "H": 12, "I": 14})
    _band(ws, "B5", "I", "ANÁLISE DE COMPARÁVEIS")
    _band(ws, "B6", "I", "Múltiplos de mercado (LTM)", fill=RED, bold=False)
    heads = ["Empresa", "Ticker", "P/L", "EV/Receita", "EV/EBITDA", "EV/EBITDA ex-IFRS16", "Margem EBITDA",
             "Valor de mercado (R$ mi)"]
    for i, h in enumerate(heads):
        cell = ws.cell(row=8, column=2 + i, value=h)
        cell.font = _font(BLACK, size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor=PINK)
        cell.border = Border(bottom=DOUBLE)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[8].height = 30
    table = comps.table if comps is not None else None
    peers = table[table["Papel"] == "Peer"] if table is not None else []
    r0 = 9
    rr = r0
    cols = [("P/L", "D", F_X), ("EV/Receita", "E", F_X), ("EV/EBITDA", "F", F_X), ("EV/EBITDA ex-IFRS16", "G", F_X),
            ("Margem EBITDA", "H", F_PCT), ("Valor de Mercado", "I", F_MM)]
    for _, prow in (peers.iterrows() if table is not None else []):
        _label(ws, f"B{rr}", prow["Empresa"])
        _label(ws, f"C{rr}", prow["Ticker"])
        for key, col, fmt in cols:
            val = _num(prow[key])
            _put(ws, f"{col}{rr}", val if val is not None else "n.m.", fmt, "hard")
        for col in "BCDEFGHI":
            ws[f"{col}{rr}"].border = Border(bottom=THIN)
        rr += 1
    r1 = rr - 1
    if table is not None:
        trow = table.iloc[0]
        _label(ws, f"B{rr}", f"Alvo: {trow['Empresa']}", bold=True)
        _label(ws, f"C{rr}", trow["Ticker"], bold=True)
        for key, col, fmt in cols:
            val = _num(trow[key])
            _put(ws, f"{col}{rr}", val if val is not None else "n.m.", fmt, "hard", bold=True)
        for col in "BCDEFGHI":
            ws[f"{col}{rr}"].fill = PatternFill("solid", fgColor="FDE9E7")
        rr += 1

    rs = rr + 1
    _band(ws, f"B{rs}", "I", "ESTATÍSTICAS DOS PEERS E PREÇO IMPLÍCITO", fill=RED, bold=False)
    stats = [("Mínimo", "MIN({rng})"), ("Percentil 25", "QUARTILE({rng},1)"), ("Mediana", "MEDIAN({rng})"),
             ("Percentil 75", "QUARTILE({rng},3)"), ("Máximo", "MAX({rng})")]
    STAT = {}
    for i, (lab_, fx) in enumerate(stats):
        row = rs + 1 + i
        _label(ws, f"B{row}", lab_, bold=True, size=10)
        STAT[lab_] = row
        for _, col, fmt in cols[:4]:
            rng = f"{col}${r0}:{col}${r1}"
            _put(ws, f"{col}{row}", f'=IF(COUNT({rng})=0,"n.m.",{fx.format(rng=rng)})', F_X, bold=True)
    # métricas do alvo usadas no preço implícito
    fin_rows = rs + 8
    _band(ws, f"B{fin_rows}", "D", "Métricas do alvo (R$ mi)", fill=RED2, size=10)
    tgt_vals = []
    if table is not None:
        t = table.iloc[0]
        tgt_vals = [("Lucro líquido LTM", t["Lucro Líquido LTM"]), ("Receita LTM", t["Receita LTM"]),
                    ("EBITDA LTM", t["EBITDA LTM"]), ("EBITDA ex-IFRS16 LTM", t["EBITDA ex-IFRS16 LTM"]),
                    ("Ponte EV → equity (dív. líq. + arrend. + minor.)", t["EV"] - t["Valor de Mercado"]),
                    ("Ponte ex-IFRS16 (dív. líq. + minor.)", t["EV ex-IFRS16"] - t["Valor de Mercado"])]
    M = {}
    for i, (lab_, val) in enumerate(tgt_vals):
        row = fin_rows + 1 + i
        _label(ws, f"B{row}", lab_, size=10)
        _put(ws, f"D{row}", float(val), F_MM, "hard")
        M[i] = f"$D${row}"
    shares_ref = "DCF!$G$8"
    ri = fin_rows + 8
    _band(ws, f"B{ri}", "G", "Preço implícito por ação (R$)", fill=RED2, size=10)
    PRICE = {}
    if tgt_vals:
        for i, (lab_, _) in enumerate(stats):
            row = ri + 1 + i
            _label(ws, f"B{row}", lab_, size=10)
            srow = STAT[lab_]
            _put(ws, f"D{row}", f'=IFERROR(D{srow}*{M[0]}/{shares_ref},"n.m.")', F_PX)
            _put(ws, f"E{row}", f'=IFERROR((E{srow}*{M[1]}-{M[4]})/{shares_ref},"n.m.")', F_PX)
            _put(ws, f"F{row}", f'=IFERROR((F{srow}*{M[2]}-{M[4]})/{shares_ref},"n.m.")', F_PX)
            _put(ws, f"G{row}", f'=IFERROR((G{srow}*{M[3]}-{M[5]})/{shares_ref},"n.m.")', F_PX)
            PRICE[lab_] = row
    _label(ws, f"B{ri + 7}", "P/L aplica o múltiplo ao lucro; múltiplos de EV aplicam ao EBITDA ou à receita e subtraem a ponte "
           "até o equity. Percentis no método QUARTILE.", color=MUTED, size=9, italic=True)

    # ================================================================== Resumo + football field
    ws = ws_r
    _frame(ws, header, "H", {"B": 40, "C": 13, "D": 13, "E": 13, "F": 15, "G": 3})
    _band(ws, "B5", "F", f"Resumo do valuation — {v.company} ({v.ticker})")
    summary_rows = [
        (7, "Recomendação", "=DCF!G15", None),
        (8, "Preço-alvo (R$, DCF)", "=DCF!G13", F_PX),
        (9, "Preço atual (R$)", "=DCF!G7", F_PX),
        (10, "Upside", "=DCF!G14", F_PCT),
        (11, "Cenário", "='Cenários'!B6", None),
        (12, "WACC", "=DCF!J38", F_PCT),
        (13, "Crescimento na perpetuidade (g)", "=DCF!J39", F_PCT),
    ]
    for row, lab_, fx, fmt in summary_rows:
        _label(ws, f"B{row}", lab_, bold=row in (7, 8))
        _put(ws, f"D{row}", fx, fmt, "link", bold=row in (7, 8), size=14 if row == 7 else 11, align="right")
    _rec_format(ws, "D7")
    ws.row_dimensions[7].height = 24
    _label(ws, "B14", f"Critério: COMPRA com upside ≥ {BUY_UPSIDE:.0%}, VENDA com upside ≤ {SELL_UPSIDE:.0%}, "
           "senão NEUTRO (editável na aba DCF).", color=MUTED, size=9, italic=True)

    _band(ws, "B16", "F", "FOOTBALL FIELD — valor por ação (R$)", fill=RED)
    for i, h in enumerate(["Método", "Mínimo", "Amplitude", "Máximo", "Ponto central"]):
        cell = ws.cell(row=17, column=2 + i, value=h)
        cell.font = _font(BLACK, size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor=PINK)
        cell.border = Border(bottom=DOUBLE)
        cell.alignment = Alignment(horizontal="center")
    rr = 18
    ff = []
    if PRICE:
        ff.append(("Comps: EV/EBITDA (P25–P75)", f"=Comps!F{PRICE['Percentil 25']}", f"=Comps!F{PRICE['Percentil 75']}",
                   f"=Comps!F{PRICE['Mediana']}", "link"))
        ff.append(("Comps: P/L (P25–P75)", f"=Comps!D{PRICE['Percentil 25']}", f"=Comps!D{PRICE['Percentil 75']}",
                   f"=Comps!D{PRICE['Mediana']}", "link"))
    if PS:
        cells = ",".join(f"'DCF Cenários'!$D${row}" for row in PS.values())
        ff.append(("DCF (pessimista–otimista)", f"=MIN({cells})", f"=MAX({cells})", "=DCF!G13", "link"))
    for row in v.field:
        if row.kind == "mercado":
            ff.append((row.label, row.low, row.high, row.mid, "hard"))
    ff.append(("Preço atual", "=DCF!G7*0.995", "=DCF!G7*1.005", "=DCF!G7", "link"))
    for lab_, lo, hi, mid, kind in ff:
        _label(ws, f"B{rr}", lab_)
        _put(ws, f"C{rr}", lo, F_PX, kind)
        _put(ws, f"E{rr}", hi, F_PX, kind)
        _put(ws, f"F{rr}", mid, F_PX, kind)
        _put(ws, f"D{rr}", f"=E{rr}-C{rr}", F_PX)
        for col in "BCDEF":
            ws[f"{col}{rr}"].border = Border(bottom=THIN)
        rr += 1
    last = rr - 1
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.title = "Valor por ação (R$)"
    ch.legend = None
    ch.add_data(Reference(ws, min_col=3, max_col=4, min_row=17, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=18, max_row=last))
    ch.series[0].graphicalProperties.noFill = True
    ch.series[0].graphicalProperties.line.noFill = True
    ch.series[1].graphicalProperties.solidFill = DARK
    ch.series[1].graphicalProperties.line.solidFill = DARK
    ch.y_axis.numFmt = '0.00'
    ch.y_axis.majorGridlines = None
    ch.x_axis.scaling.orientation = "maxMin"
    ch.height, ch.width = 7.5, 17
    ws.add_chart(ch, f"B{last + 2}")
    rt = last + 18
    _band(ws, f"B{rt}", "F", "TESES DO ANALISTA", fill=RED)
    for i, t in enumerate(v.theses or ["Sem teses informadas: premissas por regra simples a partir do histórico."]):
        c = _label(ws, f"B{rt + 1 + i}", f"{i + 1}. {t}" if v.theses else t)
        c.alignment = Alignment(wrap_text=False)
    _label(ws, f"B{rt + 2 + max(len(v.theses), 1)}",
           "Exercício acadêmico com dados públicos (CVM, Yahoo Finance, Focus, Damodaran). Não é recomendação de investimento.",
           color=MUTED, size=9, italic=True)

    # ================================================================== Sensibilidade (cenário selecionado)
    ws = ws_s
    _frame(ws, header, "I", {"B": 14, **{get_column_letter(3 + j): 12 for j in range(5)}})
    _band(ws, "B5", "G", "SENSIBILIDADE DO PREÇO-ALVO (R$) — WACC × g")
    _label(ws, "B6", "Cenário selecionado na aba Cenários. Cada célula refaz o DCF inteiro.", color=MUTED, size=9, italic=True)
    steps = [-0.01, -0.005, 0.0, 0.005, 0.01]
    _label(ws, "B8", "WACC \\ g", bold=True)
    for j, dg in enumerate(steps):
        col = get_column_letter(3 + j)
        c = _put(ws, f"{col}8", f"=DCF!$J$39+{dg}", F_PCT, "link", bold=True)
        c.fill = PatternFill("solid", fgColor=PINK)
    fc, per = "DCF!$K$19:$O$19", "DCF!$K$20:$O$20"
    for i, dw in enumerate(steps):
        row = 9 + i
        c = _put(ws, f"B{row}", f"=DCF!$J$38+{dw}", F_PCT, "link", bold=True)
        c.fill = PatternFill("solid", fgColor=PINK)
        for j in range(5):
            col = get_column_letter(3 + j)
            wv, gv = f"$B{row}", f"{col}$8"
            fx = (f"=(SUMPRODUCT({fc}/(1+{wv})^{per})+DCF!$O$19*(1+{gv})/({wv}-{gv})/(1+{wv})^5"
                  f"-DCF!$G$9-DCF!$G$10)/DCF!$G$8")
            cell = _put(ws, f"{col}{row}", fx, F_PX, gray=True)
            if i == 2 and j == 2:
                cell.fill = PatternFill("solid", fgColor=TARGET)
                cell.font = _font(BLACK, bold=True)

    # ================================================================== Teses
    ws = ws_t
    _frame(ws, header, "E", {"B": 6, "C": 110})
    _band(ws, "B5", "C", "TESES DE INVESTIMENTO E COMO VIRARAM PREMISSAS")
    _label(ws, "B6", f"Premissas: {v.assumptions_source}. A IA propõe premissas; o código valida os limites e calcula; "
           "o analista aprova.", color=MUTED, size=9, italic=True)
    for i, t in enumerate(v.theses or ["(nenhuma tese informada: premissas por regra simples)"]):
        _label(ws, f"B{8 + i}", i + 1 if v.theses else "")
        c = _label(ws, f"C{8 + i}", t)
        c.alignment = Alignment(wrap_text=True)
    rr = 10 + max(len(v.theses), 1)
    _band(ws, f"B{rr}", "C", "Avisos de validação", fill=RED, size=11)
    for i, wtxt in enumerate(v.warnings or ["Nenhum"]):
        _label(ws, f"C{rr + 1 + i}", wtxt, size=10)

    # ================================================================== Histórico
    ws = ws_h
    _frame(ws, header, "G", {"B": 36, "C": 13, "D": 13, "E": 13})
    _band(ws, "B5", "E", "HISTÓRICO (CVM) — R$ milhões")
    hist = v.history
    for j, col in enumerate(hist.columns):
        _band(ws, f"{get_column_letter(3 + j)}6", get_column_letter(3 + j), str(col), fill=RED, size=11)
    for i, (idx, row) in enumerate(hist.iterrows()):
        r_ = 7 + i
        _label(ws, f"B{r_}", idx)
        pct = ("/" in idx) or ("Margem" in idx) or ("Crescimento" in idx)
        for j, val in enumerate(row):
            if val == val:
                _put(ws, f"{get_column_letter(3 + j)}{r_}", float(val), F_PCT if pct else F_MM, "hard", gray=True)

    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToHeight = 0
    wb.save(path)
    return path
