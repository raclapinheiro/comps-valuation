"""One-pager (PowerPoint) da empresa-alvo, no estilo de pitchbook.

Todos os números vêm do CompsResult (código determinístico). A IA só fornece o texto
de descrição do negócio, que é marcado como tal no rodapé. Os gráficos são nativos do
PowerPoint (editáveis), não imagens.
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .engine import CompsResult

NAVY = RGBColor(0x1F, 0x38, 0x64)
INK = RGBColor(0x26, 0x26, 0x26)
GREY = RGBColor(0x6E, 0x6E, 0x6E)
LIGHT = RGBColor(0xEE, 0xF1, 0xF6)
LINE = RGBColor(0xC9, 0xD1, 0xDE)
ACCENT = RGBColor(0xC0, 0x1F, 0x3A)   # destaque do alvo nos gráficos
PEER = RGBColor(0x8C, 0x9B, 0xB4)
FONT = "Arial"


def _no_shadow(shape) -> None:
    """Remove a sombra herdada do tema (PowerPoint e LibreOffice)."""
    shape.shadow.inherit = False
    for ref in shape._element.xpath(".//a:effectRef"):
        ref.set("idx", "0")


def _text(slide, x, y, w, h, text, size=11, bold=False, color=INK, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    f = r.font
    f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
    f.color.rgb = color
    return box


def _section(slide, x, y, w, title):
    _text(slide, x, y, w, 0.28, title.upper(), size=10, bold=True, color=NAVY)
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y + 0.3), Inches(x + w), Inches(y + 0.3))
    ln.line.color.rgb = LINE
    ln.line.width = Pt(0.75)
    _no_shadow(ln)


def _fmt_mm(v: float) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "n.d."
    if abs(v) >= 1000:
        return f"R$ {v / 1000:,.1f} bi".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {v:,.0f} mi".replace(",", ".")


def _num(v: float, dec: int = 0) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "n.m."
    s = f"{v:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def _x(v: float) -> str:
    return "n.m." if v is None or np.isnan(v) else _num(v, 1) + "x"


def _pct(v: float) -> str:
    return "n.m." if v is None or np.isnan(v) else _num(v * 100, 1) + "%"


MESES = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]


def _br_date(iso: str) -> str:
    try:
        y, m, d = iso[:10].split("-")
        return f"{d}/{m}/{y}"
    except ValueError:
        return iso


def _style_axis_font(axis, size=8):
    axis.tick_labels.font.size = Pt(size)
    axis.tick_labels.font.name = FONT
    axis.tick_labels.font.color.rgb = GREY


def export_onepager(res: CompsResult, path: Path | str, description: list[str], description_source: str,
                    valuation=None) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    t = res.target
    row = res.table.iloc[0]
    fin = res.financials[t]
    comp = res.companies[t]
    quote = res.quotes.get(t) or res.quotes.get(comp.main_ticker) or next(
        (res.quotes[x] for x in comp.tickers.values() if res.quotes.get(x)), None)
    if quote is None or not res.target_ok:
        raise ValueError(f"{t} não tem cotação ou não entrou na tabela de comps; one-pager não gerado")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    # ------------------------------------------------------------------ cabeçalho
    _text(s, 0.5, 0.35, 9.5, 0.5, f"{row['Empresa']} ({t})", size=24, bold=True, color=NAVY)
    _text(s, 0.5, 0.88, 12.3, 0.3,
          f"Perfil da companhia  ·  {comp.sector}  ·  Fundamentos {fin.ltm_label} ({fin.basis})  ·  "
          f"Preço de fechamento em {_br_date(quote.as_of)}", size=10, color=GREY)

    # ------------------------------------------------------------------ coluna esquerda
    L, LW = 0.5, 5.9
    _section(s, L, 1.35, LW, "Visão geral do negócio")
    box = s.shapes.add_textbox(Inches(L), Inches(1.75), Inches(LW), Inches(1.95))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = 0
    for i, b in enumerate(description[:4]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(5)
        r = p.add_run()
        r.text = "▪  " + b
        r.font.name, r.font.size = FONT, Pt(11)
        r.font.color.rgb = INK

    # Indicadores de destaque
    stats = [
        ("Valor de mercado", _fmt_mm(row["Valor de Mercado"])),
        ("Enterprise value", _fmt_mm(row["EV"])),
        ("EV / EBITDA LTM", _x(row["EV/EBITDA"])),
        ("Margem EBITDA LTM", _pct(row["Margem EBITDA"])),
    ]
    cw, gap, sy = (LW - 3 * 0.15) / 4, 0.15, 3.85
    for i, (lab, val) in enumerate(stats):
        x = L + i * (cw + gap)
        card = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(sy), Inches(cw), Inches(0.95))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT
        card.line.fill.background()
        _no_shadow(card)
        _text(s, x + 0.1, sy + 0.12, cw - 0.2, 0.4, val, size=16, bold=True, color=NAVY)
        _text(s, x + 0.1, sy + 0.58, cw - 0.2, 0.3, lab, size=8.5, color=GREY)

    # Resumo financeiro
    _section(s, L, 5.0, LW, "Resumo financeiro (R$ milhões)")
    fy_prev_lbl = f"FY{int(fin.fy_label[2:]) - 1:02d}"
    cols = [fy_prev_lbl, f"FY{fin.fy_label[2:]}", fin.ltm_label]
    periods = [fin.fy_prev, fin.fy, fin.ltm]
    lines = [
        ("Receita líquida", [_num(p.revenue) for p in periods]),
        ("EBITDA", [_num(p.ebitda) for p in periods]),
        ("Margem EBITDA", [_pct(p.ebitda / p.revenue) if p.revenue else "n.d." for p in periods]),
        ("Lucro líquido (controladores)", [_num(p.net_income) for p in periods]),
        ("Dív. líq. / EBITDA (ex-IFRS16)", ["", "", _x(row["Dív. Líq./EBITDA ex-IFRS16"])]),
    ]
    tbl = s.shapes.add_table(len(lines) + 1, 4, Inches(L), Inches(5.42), Inches(LW), Inches(1.5)).table
    tbl.columns[0].width = Inches(LW - 3 * 1.05)
    for j in range(1, 4):
        tbl.columns[j].width = Inches(1.05)
    for i in range(len(lines) + 1):
        for j in range(4):
            cell = tbl.cell(i, j)
            if i == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            txt = ("" if j == 0 else cols[j - 1]) if i == 0 else (lines[i - 1][0] if j == 0 else lines[i - 1][1][j - 1])
            cell.text = txt
            if i > 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i % 2 else LIGHT
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
            for r in p.runs:
                r.font.name, r.font.size = FONT, Pt(9.5)
                r.font.bold = i == 0 or (j == 3 and i > 0)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK

    # ------------------------------------------------------------------ coluna direita
    R, RW = 6.9, 5.93
    hist = quote.history if quote is not None and quote.history is not None and len(quote.history) > 5 else None
    ret = f"  ·  {_pct(hist.iloc[-1] / hist.iloc[0] - 1)} em 12 meses" if hist is not None else ""
    _section(s, R, 1.35, RW, f"Desempenho da ação (R$){ret}")
    if hist is not None:
        weekly = hist.resample("W").last().dropna()
        cd = CategoryChartData()
        cd.categories = [f"{MESES[d.month - 1]}/{d.strftime('%y')}" for d in weekly.index]
        cd.add_series(t, [round(float(v), 2) for v in weekly.values])
        gf = s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(R), Inches(1.75), Inches(RW), Inches(2.95), cd)
        ch = gf.chart
        ch.has_title = False
        ch.has_legend = False
        ch.font.name, ch.font.size = FONT, Pt(8)
        ser = ch.series[0]
        ser.format.line.color.rgb = NAVY
        ser.format.line.width = Pt(1.75)
        ser.smooth = False
        ca, va = ch.category_axis, ch.value_axis
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        ca.format.line.color.rgb = LINE
        ca.has_major_gridlines = False
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = RGBColor(0xE4, 0xE7, 0xEC)
        va.format.line.fill.background()
        va.tick_labels.number_format = '#,##0.00'
        va.tick_labels.number_format_is_linked = False
        lo, hi = float(weekly.min()), float(weekly.max())
        pad = (hi - lo) * 0.1 or 1
        raw_step = (hi - lo + 2 * pad) / 5
        mag = 10 ** np.floor(np.log10(raw_step))
        step = next(m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= raw_step)
        va.major_unit = float(step)
        va.minimum_scale = float(max(0, np.floor((lo - pad) / step) * step))
        va.maximum_scale = float(np.ceil((hi + pad) / step) * step)
        _style_axis_font(ca)
        _style_axis_font(va)
    else:
        _text(s, R, 2.6, RW, 0.4, "Histórico de preços indisponível", size=10, color=GREY, align=PP_ALIGN.CENTER)

    # EV/EBITDA: alvo vs peers
    tb = res.table[["Ticker", "Papel", "EV/EBITDA"]].dropna().sort_values("EV/EBITDA", ascending=True)
    med = res.table[res.table["Papel"] == "Peer"]["EV/EBITDA"].median()
    _section(s, R, 4.85, RW, f"EV / EBITDA LTM vs. peers  ·  mediana dos peers {_x(med)}")
    if not tb.empty:
        cd = CategoryChartData()
        cd.categories = list(tb["Ticker"])
        cd.add_series("EV/EBITDA", [round(float(v), 1) for v in tb["EV/EBITDA"]])
        gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(R), Inches(5.22), Inches(RW), Inches(1.75), cd)
        ch = gf.chart
        ch.has_title = False
        ch.has_legend = False
        ch.font.name, ch.font.size = FONT, Pt(8)
        plot = ch.plots[0]
        plot.gap_width = 45
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format, dl.number_format_is_linked = '0.0"x"', False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size, dl.font.name = Pt(8), FONT
        dl.font.color.rgb = INK
        ser = ch.series[0]
        for i, papel in enumerate(tb["Papel"]):
            pt = ser.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = ACCENT if papel == "Alvo" else PEER
        ch.value_axis.visible = False
        ch.value_axis.has_major_gridlines = False
        ch.value_axis.minimum_scale = 0
        ch.value_axis.maximum_scale = float(tb["EV/EBITDA"].max()) * 1.18
        ca = ch.category_axis
        ca.format.line.color.rgb = LINE
        _style_axis_font(ca)

    # ------------------------------------------------------------------ rodapé
    _text(s, 0.5, 7.08, 12.3, 0.3,
          f"Fontes: CVM (DFP/ITR, dados abertos); cotações: {quote.source}. Descrição: {description_source}. "
          f"EV inclui arrendamentos (IFRS 16) e minoritários; EBITDA = EBIT + D&A. Gerado em {res.generated_at}.",
          size=7.5, color=GREY, italic=True)

    if valuation is not None:
        _valuation_slide(prs, res, valuation)
    prs.save(path)
    return path


def _valuation_slide(prs, res: CompsResult, v) -> None:
    """Slide 2: football field (gráfico nativo) + premissas por cenário + teses."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    row = res.table.iloc[0]
    _text(s, 0.5, 0.35, 9.4, 0.5, f"Valuation: {row['Empresa']} ({v.ticker})", size=24, bold=True, color=NAVY)
    rec = v.recommendation
    rec_color = {"COMPRA": RGBColor(0x1E, 0x7B, 0x34), "VENDA": RGBColor(0xB1, 0x14, 0x1A)}.get(rec, RGBColor(0x59, 0x59, 0x59))
    badge = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.95), Inches(0.3), Inches(2.88), Inches(0.82))
    badge.fill.solid()
    badge.fill.fore_color.rgb = rec_color
    badge.line.fill.background()
    _no_shadow(badge)
    _text(s, 10.05, 0.34, 2.7, 0.42, rec, size=20, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), align=PP_ALIGN.CENTER)
    _text(s, 10.05, 0.76, 2.7, 0.3, f"Alvo R$ {_num(v.target_price, 2)} · upside {_pct(v.upside)}", size=10,
          color=RGBColor(0xFF, 0xFF, 0xFF), align=PP_ALIGN.CENTER)
    _text(s, 0.5, 0.88, 12.3, 0.3,
          f"Valor por ação (R$) por método  ·  preço atual R$ {_num(v.price, 2)} em {_br_date(v.price_date)}  ·  "
          f"premissas: {v.assumptions_source}", size=10, color=GREY)

    # ---------------- football field
    L, LW = 0.5, 7.7
    _section(s, L, 1.35, LW, "Football field")
    rows = [(f"{r.label}\nR$ {_num(r.low, 2)} – {_num(r.high, 2)}", r.low, r.high, r.kind) for r in v.field]
    rows.append((f"Preço atual\nR$ {_num(v.price, 2)}", v.price * 0.995, v.price * 1.005, "preco"))
    cd = CategoryChartData()
    cd.categories = [r[0] for r in rows]
    cd.add_series("base", [round(r[1], 2) for r in rows])
    cd.add_series("faixa", [round(r[2] - r[1], 2) for r in rows])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(L), Inches(1.75), Inches(LW), Inches(5.1), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    ch.font.name, ch.font.size = FONT, Pt(9)
    plot = ch.plots[0]
    plot.gap_width = 55
    plot.overlap = 100
    inv = ch.series[0]
    inv.format.fill.background()
    inv.format.line.fill.background()
    colors = {"comps": PEER, "dcf": NAVY, "mercado": RGBColor(0xC9, 0xD1, 0xDE), "preco": ACCENT}
    for i, r in enumerate(rows):
        pt = ch.series[1].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = colors[r[3]]
    ca, va = ch.category_axis, ch.value_axis
    ca.reverse_order = True
    ca.format.line.color.rgb = LINE
    _style_axis_font(ca, 9)
    lo = min(r[1] for r in rows)
    hi = max(r[2] for r in rows)
    va.minimum_scale = max(0.0, float(np.floor(lo * 0.85)))
    va.maximum_scale = float(np.ceil(hi * 1.1))
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE4, 0xE7, 0xEC)
    va.format.line.fill.background()
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = '0.00', False
    _style_axis_font(va)

    # ---------------- premissas por cenário
    R, RW = 8.6, 4.23
    _section(s, R, 1.35, RW, "Premissas e resultado do DCF")
    names = [k for k in ("pessimista", "base", "otimista") if k in v.dcf]
    lines = [
        ("Crescimento médio", [_pct(np.mean(v.dcf[k].scenario.growth)) for k in names]),
        ("Margem EBITDA média¹", [_pct(np.mean(v.dcf[k].scenario.margin)) for k in names]),
        ("WACC", [_pct(v.dcf[k].wacc) for k in names]),
        ("g perpetuidade", [_pct(v.dcf[k].scenario.terminal_growth) for k in names]),
        ("Valor por ação (R$)", [_num(v.dcf[k].per_share, 2) for k in names]),
        ("Upside", [_pct(v.dcf[k].upside) for k in names]),
    ]
    tbl = s.shapes.add_table(len(lines) + 1, len(names) + 1, Inches(R), Inches(1.75), Inches(RW), Inches(2.3)).table
    tbl.columns[0].width = Inches(RW - 3 * 0.92)
    for j in range(1, len(names) + 1):
        tbl.columns[j].width = Inches(0.92)
    head = [""] + [k.capitalize() for k in names]
    for i in range(len(lines) + 1):
        for j in range(len(names) + 1):
            cell = tbl.cell(i, j)
            cell.text = head[j] if i == 0 else (lines[i - 1][0] if j == 0 else lines[i - 1][1][j - 1])
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (RGBColor(0xFF, 0xFF, 0xFF) if i % 2 else LIGHT)
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
            for r_ in p.runs:
                r_.font.name, r_.font.size = FONT, Pt(9)
                r_.font.bold = i == 0 or lines[i - 1][0].startswith("Valor")
                r_.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK

    _section(s, R, 4.35, RW, "Teses do analista")
    box = s.shapes.add_textbox(Inches(R), Inches(4.75), Inches(RW), Inches(2.1))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = 0
    theses = v.theses or ["Sem teses informadas: premissas por regra simples a partir do histórico."]
    for i, t in enumerate(theses[:5]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(3)
        r_ = p.add_run()
        r_.text = f"{i + 1}. {t}" if v.theses else t
        r_.font.name, r_.font.size = FONT, Pt(9.5)
        r_.font.color.rgb = INK

    _text(s, 0.5, 7.08, 12.3, 0.3,
          "¹ EBITDA ex-IFRS16 (aluguel como despesa). DCF de 5 anos, FCFF descontado pelo WACC (CAPM), valor terminal de Gordon. "
          "Comps: P25–P75 dos múltiplos LTM dos peers. Recomendação: COMPRA se upside ≥ 15%, VENDA se ≤ −5%. "
          "Exercício acadêmico, não é recomendação de investimento. Fontes: CVM, Yahoo Finance, Focus, Damodaran.",
          size=7.5, color=GREY, italic=True)
