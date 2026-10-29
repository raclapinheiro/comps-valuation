"""Exporta a tabela de comps para Excel no padrão de banco.

Convenções (padrão de modelagem em IB):
  * azul  = número digitado (input vindo da CVM / mercado)
  * preto = fórmula
  * verde = fórmula que puxa de outra aba
Todos os múltiplos, médias, medianas e o valuation implícito são FÓRMULAS vivas: se o
analista trocar um preço ou ajustar um EBITDA na aba "Dados", tudo recalcula.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .engine import CompsResult

FONT = "Arial"
NAVY = "1F3864"
BLUE, BLACK, GREEN, GREY = "0000FF", "000000", "008000", "595959"
F_MM = '#,##0;(#,##0);"-"'
F_PX = '#,##0.00;(#,##0.00);"-"'
F_SH = '#,##0.0;(#,##0.0);"-"'
F_X = '0.0"x";(0.0"x");"-"'
F_PCT = '0.0%;(0.0%);"-"'
THIN = Side(style="thin", color="BFBFBF")
MED = Side(style="medium", color=NAVY)


def _f(color=BLACK, bold=False, size=9, italic=False):
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic)


def _header(ws, row: int, labels: list[str], start_col: int = 1, height: int = 30):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = Font(name=FONT, color="FFFFFF", bold=True, size=9)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = height


# Colunas da aba Dados: (título, chave, formato, tipo) — tipo: 'in' (input) | 'f' (fórmula)
DADOS_COLS = [
    ("Ticker", "ticker", None, "in"),
    ("Empresa", "empresa", None, "in"),
    ("Papel", "papel", None, "in"),
    ("Período", "periodo", None, "in"),
    ("Base CVM", "base", None, "in"),
    ("Data do preço", "data_preco", None, "in"),
    ("Preço ON (R$)", "p_on", F_PX, "in"),
    ("Ações ON (mi)", "q_on", F_SH, "in"),
    ("Preço PN (R$)", "p_pn", F_PX, "in"),
    ("Ações PN (mi)", "q_pn", F_SH, "in"),
    ("Valor de Mercado", "mcap", F_MM, "f"),
    ("Dívida Bruta ex-Arrend.", "divida", F_MM, "in"),
    ("Caixa e Aplicações", "caixa", F_MM, "in"),
    ("Dívida Líquida", "dl", F_MM, "f"),
    ("Arrendamentos (IFRS16)", "arrend", F_MM, "in"),
    ("Minoritários", "minor", F_MM, "in"),
    ("EV", "ev", F_MM, "f"),
    ("Receita LTM", "receita", F_MM, "in"),
    ("Lucro Bruto LTM", "lb", F_MM, "in"),
    ("EBIT LTM", "ebit", F_MM, "in"),
    ("D&A LTM", "da", F_MM, "in"),
    ("EBITDA LTM", "ebitda", F_MM, "f"),
    ("Lucro Líquido LTM (controladores)", "ll", F_MM, "in"),
    ("Pagto. Arrendamento LTM", "parr", F_MM, "in"),
    ("EBITDA ex-IFRS16 LTM", "ebitda16", F_MM, "f"),
    ("EV ex-IFRS16", "ev16", F_MM, "f"),
    ("Receita FY", "rfy", F_MM, "in"),
    ("Receita FY-1", "rfy1", F_MM, "in"),
]
DC = {key: get_column_letter(i + 1) for i, (_, key, _, _) in enumerate(DADOS_COLS)}


def _dados_formula(key: str, r: int) -> str:
    c = lambda k: f"{DC[k]}{r}"  # noqa: E731
    return {
        "mcap": f"={c('p_on')}*{c('q_on')}+{c('p_pn')}*{c('q_pn')}",
        "dl": f"={c('divida')}-{c('caixa')}",
        "ev": f"={c('mcap')}+{c('dl')}+{c('arrend')}+{c('minor')}",
        "ebitda": f"={c('ebit')}+{c('da')}",
        "ebitda16": f"={c('ebitda')}-{c('parr')}",
        "ev16": f"={c('mcap')}+{c('dl')}+{c('minor')}",
    }[key]


def export_excel(res: CompsResult, path: Path | str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_c = wb.active
    ws_c.title = "Comps"
    ws_d = wb.create_sheet("Dados")
    ws_a = wb.create_sheet("Auditoria")
    ws_n = wb.create_sheet("Notas")
    for ws in (ws_c, ws_d, ws_a, ws_n):
        ws.sheet_view.showGridLines = False

    table = res.table.reset_index(drop=True)

    # ------------------------------------------------------------------ Dados
    ws_d["A1"] = "Dados de entrada (R$ milhões, exceto preços e ações)"
    ws_d["A1"].font = _f(NAVY, True, 12)
    ws_d["A2"] = "Azul = input da CVM/mercado (pode ser ajustado à mão) · Preto = fórmula · Fonte detalhada de cada número na aba Auditoria"
    ws_d["A2"].font = _f(GREY, italic=True, size=8)
    hdr = 4
    _header(ws_d, hdr, [t for t, _, _, _ in DADOS_COLS], height=42)
    dados_row = {}
    for i, row in table.iterrows():
        r = hdr + 1 + i
        t = row["Ticker"]
        dados_row[t] = r
        fin, mc = res.financials[t], res.mcaps[t]
        vals = {
            "ticker": t, "empresa": row["Empresa"], "papel": row["Papel"], "periodo": row["Período"],
            "base": row["Base"], "data_preco": row["Data preço"],
            "p_on": mc.on_price if mc.on_shares else 0, "q_on": mc.on_shares / 1e6,
            "p_pn": mc.pn_price if mc.pn_shares else 0, "q_pn": mc.pn_shares / 1e6,
            "divida": fin.debt_ex_leases, "caixa": fin.cash, "arrend": fin.leases, "minor": fin.minorities,
            "receita": fin.ltm.revenue, "lb": fin.ltm.gross_profit, "ebit": fin.ltm.ebit, "da": fin.ltm.da,
            "ll": fin.ltm.net_income, "parr": fin.ltm.lease_payments,
            "rfy": fin.fy.revenue, "rfy1": fin.fy_prev.revenue,
        }
        for j, (_, key, fmt, kind) in enumerate(DADOS_COLS):
            cell = ws_d.cell(row=r, column=j + 1)
            if kind == "f":
                cell.value = _dados_formula(key, r)
                cell.font = _f(BLACK)
            else:
                cell.value = vals[key]
                cell.font = _f(BLUE, bold=(key == "ticker"))
            if fmt:
                cell.number_format = fmt
            cell.border = Border(bottom=THIN)
        if row["Papel"] == "Alvo":
            for j in range(len(DADOS_COLS)):
                ws_d.cell(row=r, column=j + 1).fill = PatternFill("solid", fgColor="FFF2CC")
        ws_d.cell(row=r, column=8).comment = Comment(f"Escala CVM ×{mc.scale_factor}: {mc.scale_reason}. {mc.detail}", "comps-bbi")
    ws_d.column_dimensions["A"].width = 9
    ws_d.column_dimensions["B"].width = 30
    for j in range(3, len(DADOS_COLS) + 1):
        ws_d.column_dimensions[get_column_letter(j)].width = 12.5
    ws_d.freeze_panes = ws_d.cell(row=hdr + 1, column=3)

    # ------------------------------------------------------------------ Comps
    tgt = table.iloc[0]
    ws_c["A1"] = f"Comparáveis de Mercado — {tgt['Empresa']} ({tgt['Ticker']})"
    ws_c["A1"].font = _f(NAVY, True, 14)
    ws_c["A2"] = (f"R$ milhões, exceto múltiplos · Fundamentos {tgt['Período']} (CVM: DFP + ITR) · "
                  f"Preços de {tgt['Data preço']} · Gerado em {res.generated_at}")
    ws_c["A2"].font = _f(GREY, size=8, italic=True)

    cols = [
        ("Empresa", None), ("Ticker", None), ("Valor de Mercado", F_MM), ("EV", F_MM), ("Receita LTM", F_MM),
        ("EBITDA LTM", F_MM), ("Margem EBITDA", F_PCT), ("Cresc. Receita (FY)", F_PCT), ("EV / Receita", F_X),
        ("EV / EBITDA", F_X), ("EV / EBIT", F_X), ("P / L", F_X), ("EV / EBITDA ex-IFRS16", F_X),
        ("Dív. Líq. / EBITDA ex-IFRS16", F_X),
    ]
    # grupo de cabeçalho
    ws_c.merge_cells("C4:H4")
    ws_c["C4"] = "Operacional e Tamanho"
    ws_c.merge_cells("I4:N4")
    ws_c["I4"] = "Múltiplos (LTM)"
    for a in ("C4", "I4"):
        ws_c[a].font = _f(NAVY, True)
        ws_c[a].alignment = Alignment(horizontal="center")
        ws_c[a].border = Border(bottom=MED)
    hdr = 5
    _header(ws_c, hdr, [c for c, _ in cols], height=42)

    def d(key: str, r: int) -> str:
        return f"Dados!{DC[key]}{r}"

    def comps_row(r: int, dr: int, bold: bool = False):
        formulas = [
            f"={d('empresa', dr)}", f"={d('ticker', dr)}", f"={d('mcap', dr)}", f"={d('ev', dr)}",
            f"={d('receita', dr)}", f"={d('ebitda', dr)}",
            f'=IF({d("receita", dr)}>0,{d("ebitda", dr)}/{d("receita", dr)},"n.m.")',
            f'=IF({d("rfy1", dr)}>0,{d("rfy", dr)}/{d("rfy1", dr)}-1,"n.m.")',
            f'=IF({d("receita", dr)}>0,{d("ev", dr)}/{d("receita", dr)},"n.m.")',
            f'=IF({d("ebitda", dr)}>0,{d("ev", dr)}/{d("ebitda", dr)},"n.m.")',
            f'=IF({d("ebit", dr)}>0,{d("ev", dr)}/{d("ebit", dr)},"n.m.")',
            f'=IF({d("ll", dr)}>0,{d("mcap", dr)}/{d("ll", dr)},"n.m.")',
            f'=IF({d("ebitda16", dr)}>0,{d("ev16", dr)}/{d("ebitda16", dr)},"n.m.")',
            f'=IF({d("ebitda16", dr)}>0,{d("dl", dr)}/{d("ebitda16", dr)},"n.m.")',
        ]
        for j, (fx, (_, fmt)) in enumerate(zip(formulas, cols)):
            cell = ws_c.cell(row=r, column=j + 1, value=fx)
            cell.font = _f(GREEN, bold=bold)
            cell.border = Border(bottom=THIN)
            if fmt:
                cell.number_format = fmt
                cell.alignment = Alignment(horizontal="right")

    r = hdr + 1
    ws_c.cell(row=r, column=1, value="Alvo").font = _f(NAVY, True)
    r += 1
    target_r = r
    comps_row(r, dados_row[tgt["Ticker"]], bold=True)
    for j in range(len(cols)):
        ws_c.cell(row=r, column=j + 1).fill = PatternFill("solid", fgColor="FFF2CC")
    r += 2
    ws_c.cell(row=r, column=1, value="Peers").font = _f(NAVY, True)
    r += 1
    first_peer = r
    for _, row in table.iloc[1:].iterrows():
        comps_row(r, dados_row[row["Ticker"]])
        r += 1
    last_peer = r - 1

    r += 1
    stats_rows = {}
    for label, fn in (("Média", "AVERAGE"), ("Mediana", "MEDIAN"), ("Máximo", "MAX"), ("Mínimo", "MIN")):
        ws_c.cell(row=r, column=1, value=label).font = _f(BLACK, True)
        for j in range(2, len(cols)):  # da coluna C (3) em diante
            col = get_column_letter(j + 1)
            rng = f"{col}{first_peer}:{col}{last_peer}"
            cell = ws_c.cell(row=r, column=j + 1, value=f'=IF(COUNT({rng})=0,"n.m.",{fn}({rng}))')
            cell.font = _f(BLACK, bold=label in ("Média", "Mediana"))
            cell.number_format = cols[j][1]
            cell.alignment = Alignment(horizontal="right")
            if label == "Média":
                cell.border = Border(top=MED)
        if label == "Média":
            ws_c.cell(row=r, column=1).border = Border(top=MED)
            ws_c.cell(row=r, column=2).border = Border(top=MED)
        stats_rows[label] = r
        r += 1

    # --------------------------------------------- valuation implícito (EV/EBITDA dos peers)
    r += 2
    ws_c.cell(row=r, column=1, value=f"Valuation implícito de {tgt['Ticker']} — EV/EBITDA LTM dos peers").font = _f(NAVY, True, 11)
    r += 1
    _header(ws_c, r, ["", "", "P25", "Mediana", "P75"], height=18)
    col_mult = "J"
    peer_rng = f"Comps!${col_mult}${first_peer}:${col_mult}${last_peer}"
    dr = dados_row[tgt["Ticker"]]
    lines = [
        ("Múltiplo EV/EBITDA dos peers", [f"=QUARTILE({peer_rng},1)", f"=MEDIAN({peer_rng})", f"=QUARTILE({peer_rng},3)"], F_X),
        ("(×) EBITDA LTM do alvo", [f"={d('ebitda', dr)}"] * 3, F_MM),
        ("(=) EV implícito", None, F_MM),
        ("(−) Dívida líquida + arrendamentos + minoritários", [f"={d('dl', dr)}+{d('arrend', dr)}+{d('minor', dr)}"] * 3, F_MM),
        ("(=) Equity value implícito", None, F_MM),
        ("(÷) Ações totais (mi)", [f"={d('q_on', dr)}+{d('q_pn', dr)}"] * 3, F_SH),
        ("(=) Valor por ação implícito (R$)", None, F_PX),
        ("Preço atual (R$)", [f"={d('p_on', dr)}"] * 3, F_PX),
        ("Upside / (downside)", None, F_PCT),
    ]
    start = r + 1
    for k, (label, fxs, fmt) in enumerate(lines):
        rr = start + k
        ws_c.cell(row=rr, column=1, value=label).font = _f(BLACK, bold=label.startswith("(=)"))
        for j, col in enumerate("CDE"):
            if fxs is not None:
                v = fxs[j]
            elif k == 2:
                v = f"={col}{start}*{col}{start + 1}"
            elif k == 4:
                v = f"={col}{start + 2}-{col}{start + 3}"
            elif k == 6:
                v = f"=IF({col}{start + 5}>0,{col}{start + 4}/{col}{start + 5},\"n.m.\")"
            else:
                v = f"=IF(ISNUMBER({col}{start + 6}),{col}{start + 6}/{col}{start + 7}-1,\"n.m.\")"
            cell = ws_c.cell(row=rr, column=j + 3, value=v)
            cell.font = _f(GREEN if "Dados!" in v or "Comps!" in v else BLACK, bold=label.startswith("(=)"))
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal="right")
            if label.startswith("(=)"):
                cell.border = Border(top=THIN)
    r = start + len(lines) + 1

    ws_c.cell(row=r, column=1, value="Notas: EV = Valor de Mercado + Dívida Líquida + Arrendamentos (IFRS 16) + Minoritários. "
              "EBITDA = EBIT + D&A (DFC). Versões ex-IFRS16 excluem arrendamentos do EV e subtraem os pagamentos de "
              "arrendamento do EBITDA. \"n.m.\" = não significativo (denominador ≤ 0). Detalhes na aba Notas.").font = _f(GREY, size=8, italic=True)
    ws_c.column_dimensions["A"].width = 34
    ws_c.column_dimensions["B"].width = 9
    for j in range(3, len(cols) + 1):
        ws_c.column_dimensions[get_column_letter(j)].width = 12
    ws_c.freeze_panes = "C6"

    # ------------------------------------------------------------------ Auditoria
    ws_a["A1"] = "Trilha de auditoria: de onde veio cada número"
    ws_a["A1"].font = _f(NAVY, True, 12)
    _header(ws_a, 3, ["Ticker", "Métrica", "Valor (R$ mm)", "Documento CVM", "Conta(s) CVM", "Cálculo / observação", "Link do documento"], height=24)
    ar = 4
    for t, fin in res.financials.items():
        if t not in dados_row:
            continue
        mc = res.mcaps.get(t)
        q = res.quotes.get(t) or res.quotes.get(res.companies[t].main_ticker)
        extra = []
        if mc and q:
            extra.append(("Valor de mercado", mc.value_mm, f"{q.source}, {q.as_of}", "", mc.detail, ""))
        rows = [(a.metric, a.value_mm, a.document, a.accounts, a.formula, a.link) for a in fin.audit] + extra
        for metric, val, doc, acc, formula, link in rows:
            vals = [t, metric, val, doc, acc, formula, link]
            for j, v in enumerate(vals):
                cell = ws_a.cell(row=ar, column=j + 1, value=v)
                cell.font = _f(BLUE if j == 2 else BLACK)
                cell.border = Border(bottom=THIN)
                if j == 2:
                    cell.number_format = '#,##0.0;(#,##0.0);"-"'
            if link:
                c = ws_a.cell(row=ar, column=7)
                c.hyperlink = link
                c.value = "abrir na CVM"
                c.font = Font(name=FONT, size=9, color="0563C1", underline="single")
            ar += 1
    for col, w in zip("ABCDEFG", (8, 40, 13, 34, 60, 40, 14)):
        ws_a.column_dimensions[col].width = w
    ws_a.freeze_panes = "A4"
    ws_a.auto_filter.ref = f"A3:G{ar - 1}"

    # ------------------------------------------------------------------ Notas
    notes = [
        "METODOLOGIA",
        "Fonte dos fundamentos: Portal de Dados Abertos da CVM (DFP anual + ITR trimestral), versão mais recente de cada documento.",
        "Base: demonstrações consolidadas; quando a empresa só publica individual (sem controladas), usa-se a individual.",
        "LTM = último exercício (DFP) + acumulado do ano no último ITR − acumulado do mesmo período do ano anterior.",
        "EBITDA = EBIT (conta 3.05) + D&A. D&A = linhas de depreciação/amortização/exaustão do DFC (6.01.01.xx), excluindo "
        "amortização de custos de captação; se não houver, usa a DVA (7.04.01).",
        "Dívida bruta = 2.01.04 + 2.02.01 (empréstimos, financiamentos e debêntures), sem arrendamentos.",
        "Arrendamentos (IFRS 16) = contas de passivo de arrendamento identificadas no balanço, dentro ou fora de 2.01.04/2.02.01.",
        "Caixa = 1.01.01 + 1.01.02 (caixa, equivalentes e aplicações financeiras de curto prazo).",
        "EV = Valor de Mercado + Dívida Líquida + Arrendamentos + Minoritários (2.03.09). Consistente com EBITDA pós-IFRS 16.",
        "Ex-IFRS16: EV sem arrendamentos e EBITDA − pagamentos de arrendamento (principal + juros, DFC de financiamento).",
        "Valor de mercado = ações ex-tesouraria por classe (composição do capital, CVM) × preço da classe (Yahoo Finance).",
        "A CVM reporta a quantidade de ações às vezes em unidades, às vezes em milhares, sem indicar: a escala é resolvida "
        "pelo nº de ações do Yahoo ou pela plausibilidade do P/VPA (ver comentário na coluna 'Ações ON' da aba Dados).",
        "Média/mediana calculadas só sobre os peers (o alvo não entra). 'n.m.' quando o denominador é ≤ 0.",
        "",
        "LIMITAÇÕES CONHECIDAS",
        "Não ajusta itens não recorrentes (o EBITDA é o contábil reportado). Não inclui derivativos nem aplicações de longo prazo na dívida líquida.",
        "Não se aplica a bancos e seguradoras (layout de demonstrações diferente).",
        "",
        "AVISOS POR EMPRESA",
    ]
    for t, ws_list in res.warnings().items():
        for w in ws_list:
            notes.append(f"{t}: {w}")
    for i, line in enumerate(notes):
        c = ws_n.cell(row=i + 1, column=1, value=line)
        c.font = _f(NAVY, True, 10) if line.isupper() and line else _f(BLACK)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws_n.column_dimensions["A"].width = 140

    for ws in (ws_c, ws_d, ws_a, ws_n):
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToHeight = 0
    wb.save(path)
    return path
